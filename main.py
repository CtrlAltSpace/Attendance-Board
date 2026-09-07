

import sys

from pathlib import Path

from datetime import datetime, timedelta

from openpyxl import Workbook, load_workbook

from PyQt6.QtCore import QEvent, Qt
from PyQt6.QtWidgets import (
    QApplication,
    QComboBox,
    QDialog,
    QFormLayout,
    QHeaderView,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
    QPlainTextEdit,
    QSizePolicy
)


EXCEL_FILE_NAME = "attendance.xlsx"
ATTENDANCE_SHEET_NAME = "Attendance"

EXCEL_HEADERS = ["Name", "Date", "Time", "Status", "Participation", "Role", "Note"]

STUDENT_INFO_HEADERS = [
    "Name",
    "Birthdate",
    "Phone Number",
    "Email",
    "Father Name",
    "Father Number",
    "Mother Name",
    "Mother Number",
]

TABLE_HEADERS = ["Name", "Date (YYYY-MM-DD)", "Time (HH:MM:SS)", "Status", "Participation Points", "Role", "Note"]

SUMMARY_TABLE_HEADERS = [
    "Student",
    "Present",
    "Absent",
    "Late",
    "Participation Points",
    "Total",
]


class SummaryWindow(QDialog):

    def __init__(self, attendance_records, parent=None):
        super().__init__(parent)

        self.attendance_records = attendance_records

        self.setup_window()
        self.setup_widgets()
        self.setup_layout()
        self.apply_styles()

        self.update_summary_table()

    def setup_window(self):
        self.setWindowTitle("Attendance Summary")
        self.resize(720, 560)

    def setup_widgets(self):

        self.go_back_button = QPushButton("Go Back")
        self.go_back_button.setObjectName("GoBackButton")
        self.go_back_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.go_back_button.setStyleSheet(
            """
            QPushButton#GoBackButton {
                background-color: #2563eb;
                color: white;
                border: none;
                border-radius: 8px;
                padding: 10px 16px;
                font-size: 14px;
                font-weight: 700;
            }

            QPushButton#GoBackButton:hover {
                background-color: #1d4ed8;
            }

            QPushButton#GoBackButton:pressed {
                background-color: #1e40af;
            }
            """
        )
        self.go_back_button.clicked.connect(self.go_back_to_main_window)

        self.summary_title_label = QLabel("Attendance Summary")
        self.summary_title_label.setObjectName("SummaryTitleLabel")
        self.summary_title_label.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self.summary_range_label = QLabel("Summary:")
        self.summary_range_combo_box = QComboBox()
        self.summary_range_combo_box.addItems(["Last Week", "Last Month", "All Time"])

        self.summary_range_combo_box.currentTextChanged.connect(
            self.update_summary_table
        )

        self.total_records_label = QLabel()
        self.most_present_label = QLabel()
        self.most_absent_label = QLabel()
        self.most_late_label = QLabel()
        self.most_participation_label = QLabel()
        self.best_rate_label = QLabel()

        self.total_records_label.setObjectName("InsightLabel")
        self.most_present_label.setObjectName("InsightLabel")
        self.most_absent_label.setObjectName("InsightLabel")
        self.most_late_label.setObjectName("InsightLabel")
        self.most_participation_label.setObjectName("InsightLabel")
        self.best_rate_label.setObjectName("InsightLabel")

        self.summary_table = QTableWidget()
        self.summary_table.setColumnCount(len(SUMMARY_TABLE_HEADERS))
        self.summary_table.setHorizontalHeaderLabels(SUMMARY_TABLE_HEADERS)
        self.summary_table.setAlternatingRowColors(True)
        self.summary_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.summary_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.summary_table.verticalHeader().setVisible(False)
        self.summary_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Stretch
        )

    def setup_layout(self):

        main_layout = QVBoxLayout()
        main_layout.setContentsMargins(24, 22, 24, 24)
        main_layout.setSpacing(14)

        top_bar_layout = QHBoxLayout()
        top_bar_layout.addWidget(self.go_back_button)
        top_bar_layout.addStretch()

        range_layout = QHBoxLayout()
        range_layout.setSpacing(10)
        range_layout.addWidget(self.summary_range_label)
        range_layout.addWidget(self.summary_range_combo_box)

        insights_layout = QVBoxLayout()
        insights_layout.setSpacing(6)
        insights_layout.addWidget(self.total_records_label)
        insights_layout.addWidget(self.most_present_label)
        insights_layout.addWidget(self.most_absent_label)
        insights_layout.addWidget(self.most_late_label)
        insights_layout.addWidget(self.most_participation_label)
        insights_layout.addWidget(self.best_rate_label)

        main_layout.addLayout(top_bar_layout)
        main_layout.addWidget(self.summary_title_label)
        main_layout.addLayout(range_layout)
        main_layout.addLayout(insights_layout)
        main_layout.addWidget(self.summary_table)

        self.setLayout(main_layout)

    def go_back_to_main_window(self):
        main_window = self.parent()
        self.close()

        if main_window is not None:
            main_window.raise_()
            main_window.activateWindow()

    def update_summary_table(self, selected_range=None):
        range_text = selected_range or self.summary_range_combo_box.currentText()
        cutoff = None
        today = datetime.now().date()
        if range_text == "Last Week":
            cutoff = today - timedelta(days=7)
        elif range_text == "Last Month":
            cutoff = today - timedelta(days=30)

        summary = {}
        for record in self.attendance_records:
            (
                student_name,
                attendance_date,
                attendance_time,
                attendance_status,
                participation_score,
                *_,
            ) = record
            if cutoff:
                try:
                    record_date = datetime.strptime(attendance_date, "%Y-%m-%d").date()
                except ValueError:
                    continue

                if record_date < cutoff:
                    continue

            student = summary.setdefault(student_name, {
                "Present": 0,
                "Absent": 0,
                "Late": 0,
                "Participation": 0,
                "Total": 0,
            })
            if attendance_status == "Present":
                student["Present"] += 1
            elif attendance_status == "Absent":
                student["Absent"] += 1
            elif attendance_status == "Late":
                student["Late"] += 1
            student["Participation"] += participation_score
            student["Total"] += 1

        self.summary_table.setRowCount(0)
        for student_name, metrics in sorted(summary.items()):
            row_position = self.summary_table.rowCount()
            self.summary_table.insertRow(row_position)
            self.summary_table.setItem(row_position, 0, QTableWidgetItem(student_name))
            self.summary_table.setItem(row_position, 1, QTableWidgetItem(str(metrics["Present"])))
            self.summary_table.setItem(row_position, 2, QTableWidgetItem(str(metrics["Absent"])))
            self.summary_table.setItem(row_position, 3, QTableWidgetItem(str(metrics["Late"])))
            self.summary_table.setItem(row_position, 4, QTableWidgetItem(str(metrics["Participation"])))
            self.summary_table.setItem(row_position, 5, QTableWidgetItem(str(metrics["Total"])))

        total_records = sum(metrics["Total"] for metrics in summary.values())
        self.total_records_label.setText(f"Total records: {total_records}")
        self.most_present_label.setText(
            f"Most present: {max(summary.items(), key=lambda item: item[1]['Present'])[0]}"
            if summary else "Most present: N/A"
        )
        self.most_absent_label.setText(
            f"Most absent: {max(summary.items(), key=lambda item: item[1]['Absent'])[0]}"
            if summary else "Most absent: N/A"
        )
        self.most_late_label.setText(
            f"Most late: {max(summary.items(), key=lambda item: item[1]['Late'])[0]}"
            if summary else "Most late: N/A"
        )
        self.most_participation_label.setText(
            f"Most participation: {max(summary.items(), key=lambda item: item[1]['Participation'])[0]}"
            if summary else "Most participation: N/A"
        )
        self.best_rate_label.setText(
            f"Highest attendance rate: {max(summary.items(), key=lambda item: item[1]['Present'] / item[1]['Total'] if item[1]['Total'] else 0)[0]}"
            if summary else "Highest attendance rate: N/A"
        )

    def apply_styles(self):

        self.setStyleSheet(
            """
            QMainWindow {
                background-color: #f4f7fb;
        }

        QLabel {
            color: #1f2937;
        }

        QLabel#TitleLabel {
            font-size: 28px;
            font-weight: 700;
            margin-bottom: 8px;
        }

        QLineEdit,
        QComboBox {
            background-color: #ffffff;
            border: 1px solid #cbd5e1;
            border-radius: 8px;
            padding: 10px 12px;
            font-size: 14px;
            color: #111827;
        }

        QLineEdit:focus,
        QComboBox:focus {
            border: 1px solid #2563eb;
        }

        QPushButton {
            background-color: #2563eb;
            color: white;
            border: none;
            border-radius: 8px;
            padding: 12px;
            font-size: 15px;
            font-weight: 600;
        }

        QPushButton:hover {
            background-color: #1d4ed8;
        }

        QPushButton:pressed {
            background-color: #1e40af;
        }

        QPushButton#RemoveButton {
            background-color: #dc2626;
        }

        QPushButton#RemoveButton:hover {
            background-color: #b91c1c;
        }

        QPushButton#RemoveButton:pressed {
            background-color: #991b1b;
        }

        QPushButton#SummaryButton {
            background-color: #2563eb;
            color: white;
            border: none;
            border-radius: 8px;
            padding: 12px;
            font-size: 15px;
            font-weight: 600;
        }

        QPushButton#SummaryButton:hover {
            background-color: #1d4ed8;
        }

        QPushButton#SummaryButton:pressed {
            background-color: #1e40af;
        }

        QPushButton#GreenButton {
            background-color: #10b981;
            color: white;
            border: none;
            border-radius: 8px;
            padding: 12px;
            font-size: 15px;
            font-weight: 600;
        }

        QPushButton#GreenButton:hover {
            background-color: #059669;
        }

        QPushButton#GreenButton:pressed {
            background-color: #047857;
        }

        /* Specific style for participation buttons - override global styles */
        QPushButton#ParticipationButton {
            background-color: #10b981;
            color: white;
            border: none;
            border-radius: 4px;
            padding: 0px 6px;
            font-size: 12px;
            font-weight: 700;
            min-width: 36px;
            max-width: 46px;
            min-height: 22px;
            max-height: 24px;
            margin: 0px;
        }

        QPushButton#ParticipationButton:hover {
            background-color: #059669;
        }

        QPushButton#ParticipationButton:pressed {
            background-color: #047857;
        }

        QTableWidget {
            background-color: #ffffff;
            alternate-background-color: #f8fafc;
            border: 1px solid #cbd5e1;
            border-radius: 8px;
            gridline-color: #e5e7eb;
            font-size: 14px;
            color: #111827;
        }

        QHeaderView::section {
            background-color: #e2e8f0;
            color: #111827;
            padding: 9px;
            border: none;
            font-size: 14px;
            font-weight: 700;
        }

        QTableWidget::item {
            padding: 6px;
        }

        QTableWidget::item:selected {
            background-color: #bfdbfe;
            color: #111827;
        }
        """
    )


class AttendanceBoard(QMainWindow):

    def __init__(self):
        super().__init__()

        self.excel_file_path = Path(EXCEL_FILE_NAME)

        self.excel_file_ready = self.create_excel_file_if_needed()

        self.setup_window()
        self.setup_widgets()
        self.setup_layout()
        self.apply_styles()

        if self.excel_file_ready:
            self.load_attendance_records()


    def setup_window(self):
        self.setWindowTitle("Attendance Board")
        self.resize(720, 520)


    def setup_widgets(self):

        self.title_label = QLabel("Attendance Board")
        self.title_label.setObjectName("TitleLabel")
        self.title_label.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self.name_input = QLineEdit()
        self.name_input.setPlaceholderText("Enter name")

        self.status_combo_box = QComboBox()
        self.status_combo_box.addItems(["Present", "Absent", "Late"])

        self.history_label = QLabel("History")
        self.history_label.setAlignment(
            Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter
        )
        self.history_combo_box = QComboBox()

        self.history_combo_box.currentTextChanged.connect(
            self.load_attendance_records_for_selected_date
        )

        self.summary_button = QPushButton("Summary")
        self.summary_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.summary_button.clicked.connect(self.open_summary_window)

        self.add_attendance_button = QPushButton("Add Attendance")
        self.add_attendance_button.setObjectName("GreenButton")
        self.add_attendance_button.setCursor(Qt.CursorShape.PointingHandCursor)

        self.add_attendance_button.clicked.connect(self.add_attendance)

        self.edit_attendance_button = QPushButton("Update Selected")
        self.edit_attendance_button.setCursor(Qt.CursorShape.PointingHandCursor)

        self.edit_attendance_button.clicked.connect(self.edit_selected_attendance)

        self.remove_attendance_button = QPushButton("Remove Selected")
        self.remove_attendance_button.setObjectName("RemoveButton")
        self.remove_attendance_button.setCursor(Qt.CursorShape.PointingHandCursor)

        self.remove_attendance_button.clicked.connect(self.remove_selected_attendance)

        self.attendance_table = QTableWidget()
        self.attendance_table.setColumnCount(len(TABLE_HEADERS))
        self.attendance_table.setHorizontalHeaderLabels(TABLE_HEADERS)

        self.attendance_table.setAlternatingRowColors(True)

        self.attendance_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)

        self.attendance_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)

        self.attendance_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Stretch
        )

        self.attendance_table.verticalHeader().setVisible(False)

        self.attendance_table.itemSelectionChanged.connect(
            self.load_selected_attendance_into_inputs
        )

        self.attendance_table.itemDoubleClicked.connect(
            self.open_student_information_from_table
        )


    def setup_layout(self):

        central_widget = QWidget()

        main_layout = QVBoxLayout()
        main_layout.setContentsMargins(28, 24, 28, 28)
        main_layout.setSpacing(14)

        button_layout = QHBoxLayout()
        button_layout.setSpacing(10)
        button_layout.addWidget(self.add_attendance_button)
        button_layout.addWidget(self.edit_attendance_button)
        button_layout.addWidget(self.remove_attendance_button)

        history_layout = QHBoxLayout()
        history_layout.setSpacing(10)
        history_layout.addWidget(self.summary_button)
        history_layout.addWidget(self.history_label)
        history_layout.addWidget(self.history_combo_box)

        main_layout.addWidget(self.title_label)
        main_layout.addWidget(self.name_input)
        main_layout.addWidget(self.status_combo_box)
        main_layout.addLayout(button_layout)
        main_layout.addLayout(history_layout)
        main_layout.addWidget(self.attendance_table)

        central_widget.setLayout(main_layout)

        self.setCentralWidget(central_widget)


    def apply_styles(self):

        self.setStyleSheet(
            """
            QMainWindow {
                background-color: #f4f7fb;
            }

            QLabel {
                color: #1f2937;
            }

            QLabel#TitleLabel {
                font-size: 28px;
                font-weight: 700;
                margin-bottom: 8px;
            }

            QLineEdit,
            QComboBox {
                background-color: #ffffff;
                border: 1px solid #cbd5e1;
                border-radius: 8px;
                padding: 10px 12px;
                font-size: 14px;
                color: #111827;
            }

            QLineEdit:focus,
            QComboBox:focus {
                border: 1px solid #2563eb;
            }

            QPushButton {
                background-color: #2563eb;
                color: white;
                border: none;
                border-radius: 8px;
                padding: 12px;
                font-size: 15px;
                font-weight: 600;
            }

            QPushButton:hover {
                background-color: #1d4ed8;
            }

            QPushButton:pressed {
                background-color: #1e40af;
            }

            QPushButton#RemoveButton {
                background-color: #dc2626;
            }

            QPushButton#RemoveButton:hover {
                background-color: #b91c1c;
            }

            QPushButton#RemoveButton:pressed {
                background-color: #991b1b;
            }

            QPushButton#GreenButton {
                background-color: #10b981;
                color: white;
                border: none;
                border-radius: 8px;
                padding: 12px;
                font-size: 15px;
                font-weight: 600;
            }

            QPushButton#GreenButton:hover {
                background-color: #059669;
            }

            QPushButton#GreenButton:pressed {
                background-color: #047857;
            }

            QPushButton#SummaryButton {
                background-color: #2563eb;
            }

            QPushButton#SummaryButton:hover {
                background-color: #1d4ed8;
            }

            QPushButton#SummaryButton:pressed {
                background-color: #1e40af;
            }

            QTableWidget {
                background-color: #ffffff;
                alternate-background-color: #f8fafc;
                border: 1px solid #cbd5e1;
                border-radius: 8px;
                gridline-color: #e5e7eb;
                font-size: 14px;
                color: #111827;
            }

            QHeaderView::section {
                background-color: #e2e8f0;
                color: #111827;
                padding: 9px;
                border: none;
                font-size: 14px;
                font-weight: 700;
            }

            QTableWidget::item {
                padding: 6px;
            }

            QTableWidget::item:selected {
                background-color: #bfdbfe;
                color: #111827;
            }

            QTableWidget::item:hover {
                background-color: #bfdbfe;
                color: #111827;
            }
            """
        )


    def create_excel_file_if_needed(self):
        if not self.excel_file_path.exists():
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = ATTENDANCE_SHEET_NAME

            worksheet.append(EXCEL_HEADERS)

            student_info_sheet = workbook.create_sheet("StudentInfo")
            student_info_sheet.append(STUDENT_INFO_HEADERS)

            return self.save_workbook(workbook)
        else:
            if not self.update_excel_headers_if_needed():
                return False

            return self.create_student_info_sheet_if_needed()

    def save_workbook(self, workbook):
        try:
            workbook.save(self.excel_file_path)
            self.excel_file_ready = True
            return True
        except PermissionError:
            QMessageBox.critical(
                self,
                "Excel File Is Open",
                "Please close attendance.xlsx in Excel, then try again.",
            )
            return False

    def ensure_excel_file_ready(self):
        if self.excel_file_ready:
            return True

        self.excel_file_ready = self.create_excel_file_if_needed()

        if self.excel_file_ready:
            self.load_attendance_records()
            return True

        QMessageBox.warning(
            self,
            "Attendance File Not Ready",
            "attendance.xlsx is not ready yet. Please close it in Excel, then try again.",
        )
        return False

    def create_student_info_sheet_if_needed(self):
        workbook = load_workbook(self.excel_file_path)

        if "StudentInfo" not in workbook.sheetnames:
            student_info_sheet = workbook.create_sheet("StudentInfo")
            student_info_sheet.append(STUDENT_INFO_HEADERS)
            return self.save_workbook(workbook)

        student_info_sheet = workbook["StudentInfo"]
        headers_changed = False

        for column_index, header_text in enumerate(STUDENT_INFO_HEADERS, start=1):
            header_cell = student_info_sheet.cell(row=1, column=column_index)

            if header_cell.value != header_text:
                header_cell.value = header_text
                headers_changed = True

        if headers_changed:
            return self.save_workbook(workbook)

        return True

    def update_excel_headers_if_needed(self):
        workbook = load_workbook(self.excel_file_path)
        worksheet = self.get_attendance_worksheet(workbook)

        current_headers = [
            worksheet.cell(row=1, column=column_index).value
            for column_index in range(1, len(EXCEL_HEADERS) + 1)
        ]

        if current_headers == EXCEL_HEADERS:
            if getattr(workbook, "_attendance_sheet_repaired", False):
                return self.save_workbook(workbook)

            return True

        all_records = self.read_all_attendance_records()
        return self.write_all_attendance_records(all_records)

    def set_role_for_name(self, student_name, new_role):
        def _date_ge(a, b):
            try:
                return a >= b
            except Exception:
                return False

        all_records = self.read_all_attendance_records()
        changed = False
        effective_date = getattr(self, "_role_effective_date", None)

        for idx, record in enumerate(all_records):
            saved_name = record[0]
            saved_date = record[1]
            if saved_name.strip().casefold() == student_name.strip().casefold():
                if effective_date is not None:
                    if not _date_ge(saved_date, effective_date):
                        continue
                (
                    n,
                    d,
                    t,
                    s,
                    p,
                    r,
                    no,
                ) = record
                if r != new_role:
                    all_records[idx] = (n, d, t, s, p, new_role, no)
                    changed = True

        if hasattr(self, "_role_effective_date"):
            delattr(self, "_role_effective_date")

        if changed:
            if not self.write_all_attendance_records(all_records):
                return

            self.load_attendance_records_for_selected_date(
                self.history_combo_box.currentText(),
                keep_scroll_position=True,
                selected_row_to_restore=self.attendance_table.currentRow(),
            )

    def update_note_for_row(self, selected_table_row, new_note):
        if not self.ensure_excel_file_ready():
            return

        selected_date = self.history_combo_box.currentText()
        if not selected_date:
            return

        all_records = self.read_all_attendance_records()
        matching_row_counter = -1

        for record_index, record in enumerate(all_records):
            (
                student_name,
                attendance_date,
                attendance_time,
                attendance_status,
                participation_score,
                role,
                note,
            ) = record

            if attendance_date != selected_date:
                continue

            matching_row_counter += 1

            if matching_row_counter == selected_table_row:
                trimmed = new_note[:500]
                all_records[record_index] = (
                    student_name,
                    attendance_date,
                    attendance_time,
                    attendance_status,
                    participation_score,
                    role,
                    trimmed,
                )
                break

        if not self.write_all_attendance_records(all_records):
            return

        self.load_attendance_records_for_selected_date(
            selected_date,
            keep_scroll_position=True,
            selected_row_to_restore=selected_table_row,
        )

    def show_note_dialog(self, selected_table_row):
        if not self.ensure_excel_file_ready():
            return

        selected_date = self.history_combo_box.currentText()
        if not selected_date:
            return

        all_records = self.read_all_attendance_records()
        matching_row_counter = -1
        current_note = ""

        for record in all_records:
            (
                student_name,
                attendance_date,
                attendance_time,
                attendance_status,
                participation_score,
                role,
                note,
            ) = record
            if attendance_date != selected_date:
                continue
            matching_row_counter += 1
            if matching_row_counter == selected_table_row:
                current_note = note or ""
                break

        dialog = QDialog(self)
        dialog.setWindowTitle("Edit Note")
        layout = QVBoxLayout()
        editor = QPlainTextEdit()
        editor.setPlainText(current_note)
        editor.setMaximumBlockCount(1000)
        info_label = QLabel(f"Max 500 characters. Current: {len(current_note)}")

        def on_text_changed():
            txt = editor.toPlainText()
            info_label.setText(f"Max 500 characters. Current: {len(txt)}")

        editor.textChanged.connect(on_text_changed)

        btn_layout = QHBoxLayout()
        save_btn = QPushButton("Save")
        cancel_btn = QPushButton("Cancel")

        def save_and_close():
            text = editor.toPlainText()
            if len(text) > 500:
                QMessageBox.warning(dialog, "Too Long", "Note must be 500 characters or fewer.")
                return
            self.update_note_for_row(selected_table_row, text)
            dialog.accept()

        save_btn.clicked.connect(save_and_close)
        cancel_btn.clicked.connect(dialog.reject)

        btn_layout.addWidget(save_btn)
        btn_layout.addWidget(cancel_btn)

        layout.addWidget(editor)
        layout.addWidget(info_label)
        layout.addLayout(btn_layout)
        dialog.setLayout(layout)
        dialog.exec()

    def open_student_information_from_table(self, clicked_item):
        self.open_student_information_from_row(clicked_item.row())

    def enable_student_information_double_click(self, widget, row_index):
        widget.setProperty("student_info_row", row_index)
        widget.installEventFilter(self)

    def eventFilter(self, watched_object, event):
        if event.type() == QEvent.Type.MouseButtonDblClick:
            row_index = watched_object.property("student_info_row")

            if row_index is not None:
                self.open_student_information_from_row(int(row_index))
                return True

        return super().eventFilter(watched_object, event)

    def open_student_information_from_row(self, row_index):
        name_item = self.attendance_table.item(row_index, 0)

        if name_item is None:
            return

        student_name = name_item.text().strip()

        if not student_name:
            return

        self.open_student_information_dialog(student_name)

    def open_student_information_dialog(self, student_name):
        if not self.ensure_excel_file_ready():
            return

        student_information = self.load_student_information(student_name)

        dialog = QDialog(self)
        dialog.setWindowTitle(f"Student Information - {student_name}")
        dialog.resize(480, 360)

        main_layout = QVBoxLayout()
        main_layout.setContentsMargins(24, 22, 24, 24)
        main_layout.setSpacing(14)

        title_label = QLabel(f"Information: {student_name}")
        title_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        title_label.setStyleSheet(
            "font-size: 20px; font-weight: 700; color: #111827;"
        )

        form_layout = QFormLayout()
        form_layout.setSpacing(10)

        birthdate_input = QLineEdit(student_information["Birthdate"])
        phone_number_input = QLineEdit(student_information["Phone Number"])
        email_input = QLineEdit(student_information["Email"])
        father_name_input = QLineEdit(student_information["Father Name"])
        father_number_input = QLineEdit(student_information["Father Number"])
        mother_name_input = QLineEdit(student_information["Mother Name"])
        mother_number_input = QLineEdit(student_information["Mother Number"])

        birthdate_input.setPlaceholderText("YYYY-MM-DD")
        phone_number_input.setPlaceholderText("Student phone number")
        email_input.setPlaceholderText("Student email")
        father_name_input.setPlaceholderText("Father's name")
        father_number_input.setPlaceholderText("Father's phone number")
        mother_name_input.setPlaceholderText("Mother's name")
        mother_number_input.setPlaceholderText("Mother's phone number")

        form_layout.addRow("Birthdate:", birthdate_input)
        form_layout.addRow("Phone Number:", phone_number_input)
        form_layout.addRow("Email:", email_input)
        form_layout.addRow("Father's Name:", father_name_input)
        form_layout.addRow("Father's Number:", father_number_input)
        form_layout.addRow("Mother's Name:", mother_name_input)
        form_layout.addRow("Mother's Number:", mother_number_input)

        button_layout = QHBoxLayout()
        button_layout.addStretch()

        save_button = QPushButton("Save")
        save_button.setObjectName("GreenButton")
        save_button.setCursor(Qt.CursorShape.PointingHandCursor)

        cancel_button = QPushButton("Cancel")
        cancel_button.setCursor(Qt.CursorShape.PointingHandCursor)

        def save_information_and_close():
            updated_information = {
                "Birthdate": birthdate_input.text().strip(),
                "Phone Number": phone_number_input.text().strip(),
                "Email": email_input.text().strip(),
                "Father Name": father_name_input.text().strip(),
                "Father Number": father_number_input.text().strip(),
                "Mother Name": mother_name_input.text().strip(),
                "Mother Number": mother_number_input.text().strip(),
            }

            if not self.save_student_information(student_name, updated_information):
                return

            QMessageBox.information(
                dialog,
                "Information Saved",
                "Student information has been saved successfully.",
            )
            dialog.accept()

        save_button.clicked.connect(save_information_and_close)
        cancel_button.clicked.connect(dialog.reject)

        button_layout.addWidget(save_button)
        button_layout.addWidget(cancel_button)

        dialog.setStyleSheet(
            """
            QDialog {
                background-color: #f4f7fb;
            }

            QLabel {
                color: #111827;
                font-size: 14px;
            }

            QLineEdit {
                background-color: #ffffff;
                border: 1px solid #cbd5e1;
                border-radius: 8px;
                padding: 9px 12px;
                font-size: 14px;
                color: #111827;
            }

            QLineEdit:focus {
                border: 1px solid #2563eb;
            }

            QPushButton {
                background-color: #2563eb;
                color: white;
                border: none;
                border-radius: 8px;
                padding: 10px 16px;
                font-weight: 700;
            }

            QPushButton:hover {
                background-color: #1d4ed8;
            }

            QPushButton:pressed {
                background-color: #1e40af;
            }

            QPushButton#GreenButton {
                background-color: #10b981;
            }

            QPushButton#GreenButton:hover {
                background-color: #059669;
            }

            QPushButton#GreenButton:pressed {
                background-color: #047857;
            }
            """
        )

        main_layout.addWidget(title_label)
        main_layout.addLayout(form_layout)
        main_layout.addLayout(button_layout)
        dialog.setLayout(main_layout)
        dialog.exec()

    def load_student_information(self, student_name):
        blank_information = {
            "Birthdate": "",
            "Phone Number": "",
            "Email": "",
            "Father Name": "",
            "Father Number": "",
            "Mother Name": "",
            "Mother Number": "",
        }

        if not self.create_student_info_sheet_if_needed():
            return blank_information

        workbook = load_workbook(self.excel_file_path)
        student_info_sheet = workbook["StudentInfo"]
        header_values = [
            student_info_sheet.cell(row=1, column=column_index).value
            for column_index in range(1, student_info_sheet.max_column + 1)
        ]
        cleaned_student_name = student_name.strip().casefold()

        def get_student_info_value(row, header_name, old_header_name, fallback_index):
            for possible_header in [header_name, old_header_name]:
                if possible_header in header_values:
                    header_index = header_values.index(possible_header)

                    if header_index < len(row) and row[header_index] is not None:
                        return str(row[header_index])

            if fallback_index < len(row) and row[fallback_index] is not None:
                return str(row[fallback_index])

            return ""

        for row in student_info_sheet.iter_rows(min_row=2, values_only=True):
            saved_name = row[0] if len(row) > 0 else None

            if saved_name is None:
                continue

            if str(saved_name).strip().casefold() != cleaned_student_name:
                continue

            return {
                "Birthdate": get_student_info_value(row, "Birthdate", "Birthdate", 1),
                "Phone Number": get_student_info_value(row, "Phone Number", "Phone Number", 2),
                "Email": get_student_info_value(row, "Email", "Email", 3),
                "Father Name": get_student_info_value(row, "Father Name", "Parent Name", 4),
                "Father Number": get_student_info_value(row, "Father Number", "Parent Number", 5),
                "Mother Name": get_student_info_value(row, "Mother Name", None, 6),
                "Mother Number": get_student_info_value(row, "Mother Number", None, 7),
            }

        return blank_information

    def save_student_information(self, student_name, student_information):
        if not self.create_student_info_sheet_if_needed():
            return False

        workbook = load_workbook(self.excel_file_path)
        student_info_sheet = workbook["StudentInfo"]
        cleaned_student_name = student_name.strip().casefold()
        matching_row_index = None

        for row_index in range(2, student_info_sheet.max_row + 1):
            saved_name = student_info_sheet.cell(row=row_index, column=1).value

            if saved_name is None:
                continue

            if str(saved_name).strip().casefold() == cleaned_student_name:
                matching_row_index = row_index
                break

        if matching_row_index is None:
            matching_row_index = student_info_sheet.max_row + 1

        values_to_save = [
            student_name,
            student_information["Birthdate"],
            student_information["Phone Number"],
            student_information["Email"],
            student_information["Father Name"],
            student_information["Father Number"],
            student_information["Mother Name"],
            student_information["Mother Number"],
        ]

        for column_index, value in enumerate(values_to_save, start=1):
            student_info_sheet.cell(
                row=matching_row_index,
                column=column_index,
            ).value = value

        return self.save_workbook(workbook)


    def load_attendance_records(self):
        self.refresh_history_dropdown()

    def open_summary_window(self):
        if not self.ensure_excel_file_ready():
            return

        all_records = self.read_all_attendance_records()
        summary_window = SummaryWindow(all_records, self)
        summary_window.showMaximized()
        summary_window.exec()

    def refresh_history_dropdown(self, selected_date=None):

        if selected_date is None:
            selected_date = self.get_today_date_string()

        all_records = self.read_all_attendance_records()
        attendance_dates = {record[1] for record in all_records}
        attendance_dates.add(self.get_today_date_string())

        sorted_dates = sorted(attendance_dates, reverse=True)

        self.history_combo_box.blockSignals(True)
        self.history_combo_box.clear()
        self.history_combo_box.addItems(sorted_dates)

        selected_date_index = self.history_combo_box.findText(selected_date)

        if selected_date_index >= 0:
            self.history_combo_box.setCurrentIndex(selected_date_index)
        else:
            self.history_combo_box.setCurrentIndex(0)

        self.history_combo_box.blockSignals(False)
        self.load_attendance_records_for_selected_date()

    def load_attendance_records_for_selected_date(
        self,
        selected_date=None,
        keep_scroll_position=False,
        selected_row_to_restore=None,
    ):

        if selected_date is None:
            selected_date = self.history_combo_box.currentText()

        if not selected_date:
            return

        scroll_bar = self.attendance_table.verticalScrollBar()
        old_scroll_position = scroll_bar.value()

        self.attendance_table.setRowCount(0)
        self.name_input.clear()

        all_records = self.read_all_attendance_records()

        for (
            student_name,
            attendance_date,
            attendance_time,
            attendance_status,
            participation_score,
            role,
            note,
        ) in all_records:
            if attendance_date == selected_date:
                self.add_row_to_table(
                    student_name=student_name,
                    attendance_date=attendance_date,
                    attendance_time=attendance_time,
                    attendance_status=attendance_status,
                    participation_score=participation_score,
                    role=role,
                    note=note,
                )

        self.attendance_table.resizeColumnsToContents()
        self.attendance_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Stretch
        )

        if selected_row_to_restore is not None:
            if 0 <= selected_row_to_restore < self.attendance_table.rowCount():
                self.attendance_table.selectRow(selected_row_to_restore)

        if keep_scroll_position:
            scroll_bar.setValue(old_scroll_position)


    def add_attendance(self):

        student_name = self.name_input.text().strip()

        if not student_name:
            QMessageBox.warning(
                self,
                "Missing Name",
                "Please enter a name before adding attendance.",
            )
            self.name_input.setFocus()
            return

        if not self.ensure_excel_file_ready():
            return

        attendance_status = self.status_combo_box.currentText()

        attendance_date = self.get_today_date_string()

        attendance_time = self.get_current_time_string()

        if self.attendance_already_exists(student_name, attendance_date):
            QMessageBox.warning(
                self,
                "Duplicate Attendance",
                "This student already has an attendance record for today.",
            )
            self.name_input.setFocus()
            return

        default_role = self.get_default_role_for_name(student_name)
        if not self.save_attendance_to_excel(
            student_name,
            attendance_date,
            attendance_time,
            attendance_status,
            participation_score=0,
            role=default_role,
            note="",
        ):
            return

        self.refresh_history_dropdown(selected_date=attendance_date)

        self.name_input.clear()
        self.name_input.setFocus()

        QMessageBox.information(
            self,
            "Attendance Added",
            "The attendance record has been added successfully.",
        )

    def load_selected_attendance_into_inputs(self):

        selected_items = self.attendance_table.selectedItems()

        if not selected_items:
            return

        selected_row = self.attendance_table.currentRow()

        name_item = self.attendance_table.item(selected_row, 0)
        status_item = self.attendance_table.item(selected_row, 3)

        if name_item is None or status_item is None:
            return

        self.name_input.setText(name_item.text())
        self.status_combo_box.setCurrentText(status_item.text())

    def edit_selected_attendance(self):
        if not self.ensure_excel_file_ready():
            return

        selected_row = self.attendance_table.currentRow()

        if selected_row < 0:
            QMessageBox.warning(
                self,
                "No Row Selected",
                "Please select an attendance row to edit.",
            )
            return

        updated_student_name = self.name_input.text().strip()

        if not updated_student_name:
            QMessageBox.warning(
                self,
                "Missing Name",
                "Please enter a student name before updating attendance.",
            )
            self.name_input.setFocus()
            return

        updated_attendance_status = self.status_combo_box.currentText()

        selected_date = self.history_combo_box.currentText()

        if self.attendance_already_exists_except_row(
            updated_student_name,
            selected_date,
            selected_row,
        ):
            QMessageBox.warning(
                self,
                "Duplicate Attendance",
                "This student already has an attendance record for this date.",
            )
            self.name_input.setFocus()
            return

        if not self.update_attendance_record_for_date(
            selected_date=selected_date,
            selected_table_row=selected_row,
            updated_student_name=updated_student_name,
            updated_attendance_status=updated_attendance_status,
        ):
            return

        self.load_attendance_records_for_selected_date(
            selected_date,
            keep_scroll_position=True,
            selected_row_to_restore=selected_row,
        )

        self.attendance_table.resizeColumnsToContents()
        self.attendance_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Stretch
        )

        QMessageBox.information(
            self,
            "Attendance Updated",
            "The selected attendance record has been updated successfully.",
        )

    def remove_selected_attendance(self):
        if not self.ensure_excel_file_ready():
            return

        selected_row = self.attendance_table.currentRow()

        if selected_row < 0:
            QMessageBox.warning(
                self,
                "No Row Selected",
                "Please select an attendance row to remove.",
            )
            return

        confirmation_result = QMessageBox.question(
            self,
            "Remove Attendance",
            "Are you sure you want to remove the selected attendance record?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )

        if confirmation_result != QMessageBox.StandardButton.Yes:
            return

        selected_date = self.history_combo_box.currentText()

        if not self.remove_attendance_record_for_date(
            selected_date=selected_date,
            selected_table_row=selected_row,
        ):
            return

        self.refresh_history_dropdown(selected_date=selected_date)

        self.name_input.clear()
        self.name_input.setFocus()

        QMessageBox.information(
            self,
            "Attendance Removed",
            "The selected attendance record has been removed successfully.",
        )

    def add_row_to_table(
        self,
        student_name,
        attendance_date,
        attendance_time,
        attendance_status,
        participation_score,
        role="Member",
        note="",
    ):
        new_row_position = self.attendance_table.rowCount()
        
        self.attendance_table.insertRow(new_row_position)
        
        name_item = QTableWidgetItem(student_name)
        date_item = QTableWidgetItem(attendance_date)
        time_item = QTableWidgetItem(attendance_time)
        status_item = QTableWidgetItem(attendance_status)
        
        date_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
        time_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
        status_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
        
        self.attendance_table.setItem(new_row_position, 0, name_item)
        self.attendance_table.setItem(new_row_position, 1, date_item)
        self.attendance_table.setItem(new_row_position, 2, time_item)
        self.attendance_table.setItem(new_row_position, 3, status_item)
        
        participation_widget = QWidget()
        participation_layout = QHBoxLayout()
        participation_layout.setContentsMargins(2, 0, 2, 0)
        participation_layout.setSpacing(2)
        
        participation_score_label = QLabel(str(participation_score))
        participation_score_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        participation_score_label.setMinimumWidth(18)
        participation_score_label.setMaximumWidth(18)
        participation_score_label.setContentsMargins(0, 0, 0, 0)
        
        add_point_button = QPushButton("+1")
        add_point_button.setObjectName("ParticipationButton")
        add_point_button.setCursor(Qt.CursorShape.PointingHandCursor)
        add_point_button.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        add_point_button.setMinimumSize(42, 20)
        add_point_button.setFixedSize(42, 20)
        add_point_button.setStyleSheet(
            "background-color: #10b981; color: white; border: none; border-radius: 4px; padding: 0px 6px; font-size: 12px; font-weight: 700;"
        )
        add_point_button.clicked.connect(
            lambda checked=False, row=new_row_position: self.add_participation_point(row)
        )
        
        participation_layout.addWidget(participation_score_label)
        participation_layout.addWidget(add_point_button)
        participation_widget.setLayout(participation_layout)
        self.enable_student_information_double_click(
            participation_widget,
            new_row_position,
        )
        self.enable_student_information_double_click(
            participation_score_label,
            new_row_position,
        )
        
        self.attendance_table.setCellWidget(new_row_position, 4, participation_widget)

        role_combo = QComboBox()
        role_combo.addItems(["Member", "Staff"])
        role_combo.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        role_combo.setMaximumWidth(90)
        role_combo.setEditable(False)
        role_combo.setStyleSheet(
            "QComboBox { color: #111827; background-color: #ffffff; padding: 2px 6px; } QComboBox::drop-down { subcontrol-origin: padding; subcontrol-position: top right; width: 18px; }"
        )
        role_combo.blockSignals(True)
        role_combo.setCurrentText(role)
        role_combo.blockSignals(False)
        eff_date = self.history_combo_box.currentText()
        def _on_role_change(new_role, name=student_name, eff=eff_date):
            setattr(self, "_role_effective_date", eff)
            self.set_role_for_name(name, new_role)

        role_combo.currentTextChanged.connect(_on_role_change)
        role_widget = QWidget()
        role_layout = QHBoxLayout()
        role_layout.setContentsMargins(0, 0, 0, 0)
        role_layout.addWidget(role_combo)
        role_layout.addStretch()
        role_widget.setLayout(role_layout)
        self.enable_student_information_double_click(role_widget, new_row_position)
        self.attendance_table.setCellWidget(new_row_position, 5, role_widget)

        note_widget = QWidget()
        note_layout = QHBoxLayout()
        note_layout.setContentsMargins(0, 0, 0, 0)
        note_layout.setSpacing(6)

        first_note_line = note.splitlines()[0] if note.splitlines() else ""
        truncated = first_note_line if len(first_note_line) <= 30 else first_note_line[:27] + "..."
        note_label = QLabel(truncated)
        note_label.setObjectName("NotePreview")
        note_label.setStyleSheet("color: #111827;")
        more_button = QPushButton("More...")
        more_button.setCursor(Qt.CursorShape.PointingHandCursor)
        more_button.setMaximumWidth(80)
        more_button.setStyleSheet(
            "background-color: #2563eb; color: white; border: none; border-radius: 4px; padding: 2px 6px;"
        )
        more_button.clicked.connect(
            lambda checked=False, row=new_row_position: self.show_note_dialog(row)
        )
        note_layout.addWidget(note_label)
        note_layout.addWidget(more_button)
        note_widget.setLayout(note_layout)
        self.enable_student_information_double_click(note_widget, new_row_position)
        self.enable_student_information_double_click(note_label, new_row_position)
        self.attendance_table.setCellWidget(new_row_position, 6, note_widget)
        
        self.attendance_table.setRowHeight(new_row_position, 32)
        
        self.attendance_table.resizeColumnsToContents()
        self.attendance_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Stretch
        )


    def save_attendance_to_excel(
        self,
        student_name,
        attendance_date,
        attendance_time,
        attendance_status,
        participation_score,
        role="Member",
        note="",
    ):
        workbook = load_workbook(self.excel_file_path)
        worksheet = self.get_attendance_worksheet(workbook)

        worksheet.append(
            [
                student_name,
                attendance_date,
                attendance_time,
                attendance_status,
                participation_score,
                role,
                note,
            ]
        )

        return self.save_workbook(workbook)

    def add_participation_point(self, selected_table_row):
        if not self.ensure_excel_file_ready():
            return

        selected_date = self.history_combo_box.currentText()

        if not selected_date:
            return

        all_records = self.read_all_attendance_records()
        matching_row_counter = -1

        for record_index, record in enumerate(all_records):
            (
                student_name,
                attendance_date,
                attendance_time,
                attendance_status,
                participation_score,
                role,
                note,
            ) = record

            if attendance_date != selected_date:
                continue

            matching_row_counter += 1

            if matching_row_counter == selected_table_row:
                all_records[record_index] = (
                    student_name,
                    attendance_date,
                    attendance_time,
                    attendance_status,
                    participation_score + 1,
                    role,
                    note,
                )
                break

        if not self.write_all_attendance_records(all_records):
            return

        self.load_attendance_records_for_selected_date(
            selected_date,
            keep_scroll_position=True,
            selected_row_to_restore=selected_table_row,
        )

    def get_today_date_string(self):
        return datetime.now().strftime("%Y-%m-%d")

    def get_current_time_string(self):
        return datetime.now().strftime("%H:%M:%S")

    def normalize_date_value(self, date_value):
        if hasattr(date_value, "strftime"):
            return date_value.strftime("%Y-%m-%d")

        return str(date_value)

    def read_all_attendance_records(self):
        workbook = load_workbook(self.excel_file_path)
        worksheet = self.get_attendance_worksheet(workbook)

        if getattr(workbook, "_attendance_sheet_repaired", False):
            self.save_workbook(workbook)

        all_records = []
        header_values = [
            worksheet.cell(row=1, column=column_index).value
            for column_index in range(1, worksheet.max_column + 1)
        ]

        def get_value(row_values, header_name, fallback_index, default_value=""):
            if header_name in header_values:
                header_index = header_values.index(header_name)

                if header_index < len(row_values):
                    value = row_values[header_index]
                    if value is not None:
                        return value

            if fallback_index < len(row_values):
                value = row_values[fallback_index]
                if value is not None:
                    return value

            return default_value

        file_has_time_column = "Time" in header_values

        for row in worksheet.iter_rows(min_row=2, values_only=True):
            row_values = list(row)
            name = get_value(row_values, "Name", 0, None)
            date = get_value(row_values, "Date", 1, None)

            if file_has_time_column:
                time = get_value(row_values, "Time", 2, "")
                status = get_value(row_values, "Status", 3, None)
                participation_score = get_value(row_values, "Participation", 4, 0)
                role = get_value(row_values, "Role", 5, "Member")
                note = get_value(row_values, "Note", 6, "")
            else:
                time = ""
                status = get_value(row_values, "Status", 2, None)
                participation_score = get_value(row_values, "Participation", 3, 0)
                role = get_value(row_values, "Role", 4, "Member")
                note = get_value(row_values, "Note", 5, "")

            if name is None and date is None and status is None:
                continue

            all_records.append(
                (
                    str(name),
                    self.normalize_date_value(date),
                    self.normalize_time_value(time),
                    str(status),
                    self.normalize_participation_score(participation_score),
                    str(role),
                    str(note),
                )
            )

        return all_records

    def normalize_time_value(self, time_value):
        if time_value is None:
            return ""

        if hasattr(time_value, "strftime"):
            return time_value.strftime("%H:%M:%S")

        return str(time_value)

    def normalize_participation_score(self, participation_score):
        try:
            return int(participation_score)
        except (TypeError, ValueError):
            return 0

    def get_default_role_for_name(self, student_name):
        cleaned = student_name.strip().casefold()
        all_records = self.read_all_attendance_records()
        for record in reversed(all_records):
            saved_name = record[0]
            saved_role = record[5] if len(record) > 5 else "Member"
            if saved_name.strip().casefold() == cleaned:
                return saved_role
        return "Member"

    def attendance_already_exists(self, student_name, attendance_date):
        cleaned_student_name = student_name.strip().casefold()
        all_records = self.read_all_attendance_records()

        for (
            saved_name,
            saved_date,
            saved_time,
            saved_status,
            saved_participation_score,
            saved_role,
            saved_note,
        ) in all_records:
            if saved_date != attendance_date:
                continue

            if saved_name.strip().casefold() == cleaned_student_name:
                return True

        return False

    def attendance_already_exists_except_row(
        self,
        student_name,
        attendance_date,
        excluded_table_row,
    ):
        cleaned_student_name = student_name.strip().casefold()
        all_records = self.read_all_attendance_records()
        matching_row_counter = -1

        for (
            saved_name,
            saved_date,
            saved_time,
            saved_status,
            saved_participation_score,
            saved_role,
            saved_note,
        ) in all_records:
            if saved_date != attendance_date:
                continue

            matching_row_counter += 1

            if matching_row_counter == excluded_table_row:
                continue

            if saved_name.strip().casefold() == cleaned_student_name:
                return True

        return False

    def get_attendance_worksheet(self, workbook):
        if ATTENDANCE_SHEET_NAME in workbook.sheetnames:
            return workbook[ATTENDANCE_SHEET_NAME]

        for worksheet in workbook.worksheets:
            header_values = [
                worksheet.cell(row=1, column=column_index).value
                for column_index in range(1, worksheet.max_column + 1)
            ]
            header_set = {str(header).strip() for header in header_values if header is not None}

            if {"Name", "Date", "Status"}.issubset(header_set):
                setattr(workbook, "_attendance_sheet_repaired", True)
                worksheet.title = ATTENDANCE_SHEET_NAME
                return worksheet

        worksheet = workbook.active

        if worksheet.title != "StudentInfo":
            setattr(workbook, "_attendance_sheet_repaired", True)
            worksheet.title = ATTENDANCE_SHEET_NAME
            return worksheet

        setattr(workbook, "_attendance_sheet_repaired", True)
        worksheet = workbook.create_sheet(ATTENDANCE_SHEET_NAME, 0)
        worksheet.append(EXCEL_HEADERS)
        return worksheet

    def write_all_attendance_records(self, all_records):

        if self.excel_file_path.exists():
            workbook = load_workbook(self.excel_file_path)
            worksheet = self.get_attendance_worksheet(workbook)
            worksheet.delete_rows(1, worksheet.max_row)
        else:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = ATTENDANCE_SHEET_NAME

        worksheet.append(EXCEL_HEADERS)

        for (
            student_name,
            attendance_date,
            attendance_time,
            attendance_status,
            participation_score,
            role,
            note,
        ) in all_records:
            worksheet.append(
                [
                    student_name,
                    attendance_date,
                    attendance_time,
                    attendance_status,
                    participation_score,
                    role,
                    note,
                ]
            )

        return self.save_workbook(workbook)

    def update_attendance_record_for_date(
        self,
        selected_date,
        selected_table_row,
        updated_student_name,
        updated_attendance_status,
    ):
        all_records = self.read_all_attendance_records()
        matching_row_counter = -1

        for record_index, record in enumerate(all_records):
            (
                student_name,
                attendance_date,
                attendance_time,
                attendance_status,
                participation_score,
                role,
                note,
            ) = record

            if attendance_date != selected_date:
                continue

            matching_row_counter += 1

            if matching_row_counter == selected_table_row:
                all_records[record_index] = (
                    updated_student_name,
                    attendance_date,
                    attendance_time,
                    updated_attendance_status,
                    participation_score,
                    role,
                    note,
                )
                break

        return self.write_all_attendance_records(all_records)

    def remove_attendance_record_for_date(self, selected_date, selected_table_row):
        all_records = self.read_all_attendance_records()
        matching_row_counter = -1
        record_index_to_remove = None

        for record_index, record in enumerate(all_records):
            (
                student_name,
                attendance_date,
                attendance_time,
                attendance_status,
                participation_score,
                role,
                note,
            ) = record

            if attendance_date != selected_date:
                continue

            matching_row_counter += 1

            if matching_row_counter == selected_table_row:
                record_index_to_remove = record_index
                break

        if record_index_to_remove is not None:
            all_records.pop(record_index_to_remove)

        return self.write_all_attendance_records(all_records)


def main():

    app = QApplication(sys.argv)

    window = AttendanceBoard()

    window.showMaximized()

    sys.exit(app.exec())


if __name__ == "__main__":
    main()
